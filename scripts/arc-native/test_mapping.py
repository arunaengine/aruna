"""Checks ISA conversion, metadata preservation and malformed ARC rejection."""
# Copyright (c) 2026 The Aruna Contributors
# SPDX-License-Identifier: MIT or Apache-2.0

import base64
import datetime
import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from arctrl import ARC, ArcStudy, CompositeCell, CompositeHeader, OntologyAnnotation, Person
from openpyxl import load_workbook
from schema_salad.exceptions import ValidationException
from arc import scaffold

spec = importlib.util.spec_from_file_location("aruna_conversion", Path(__file__).resolve().parents[2] / "blob/src/arc.py")
conversion = importlib.util.module_from_spec(spec)
spec.loader.exec_module(conversion)


def source():
    return {"@context": "https://w3id.org/ro/crate/1.3/context", "@graph": [
        {"@id": "ro-crate-metadata.json", "@type": "CreativeWork",
         "about": {"@id": "urn:aruna:source"}, "conformsTo": {"@id": "https://w3id.org/ro/crate/1.3"}},
        {"@id": "urn:aruna:source", "@type": "Dataset", "identifier": "original-identifier",
         "name": "Supplied title", "description": "Supplied description", "datePublished": "2026-09-22",
         "https://example.org/custom": {"@value": "unaltered extension", "@language": "en"}},
        {"@id": "#context", "@type": "Thing", "name": "Unmapped contextual entity"},
    ]}


def files(root):
    return {str(path.relative_to(root)): base64.b64encode(path.read_bytes()).decode()
            for path in root.rglob("*") if path.is_file()}


class MappingTests(unittest.TestCase):
    def test_generation_reproducible(self):
        request = {"mode": "generate", "document_id": "stable-id", "jsonld": json.dumps(source())}
        results = []
        for year in (2001, 2031):
            class Clock(datetime.datetime):
                @classmethod
                def now(cls, tz=None):
                    return cls(year, 2, 3, 4, 5, 6, tzinfo=tz)

            with patch("datetime.datetime", Clock), patch("zipfile.time.localtime", return_value=(year, 2, 3, 4, 5, 6, 0, 1, 0)):
                results.append(conversion.convert(request))
        self.assertEqual(results[0], results[1])
        parsed = conversion.convert({"mode": "inspect", "files": results[0]["files"]})
        root = next(item for item in parsed["rocrate"]["@graph"] if item.get("@id") == "./")
        self.assertEqual(root["name"], "Supplied title")

    def test_minimal_preserved(self):
        original = json.dumps(source())
        result = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": original})
        self.assertEqual(base64.b64decode(result["files"]["aruna-metadata.json"]).decode(), original)
        self.assertFalse(any(name.endswith(("isa.study.xlsx", "isa.assay.xlsx")) for name in result["files"]))
        derived = conversion.convert({"mode": "inspect", "files": result["files"]})["rocrate"]
        root = next(item for item in derived["@graph"] if item.get("@id") == "./")
        self.assertEqual(root["name"], "Supplied title")
        self.assertEqual(root["description"], "Supplied description")
        self.assertEqual(root["identifier"], "original-identifier")
        self.assertEqual(root["datePublished"], "2026-09-22")
        self.assertEqual(root["additionalType"], "Investigation")

    def test_license_preserved(self):
        uri = "https://creativecommons.org/licenses/by/4.0/"
        for license in [uri, {"@id": uri}]:
            document = source()
            document["@graph"][1]["license"] = license
            with self.subTest(license=license):
                result = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(document)})
                self.assertEqual(base64.b64decode(result["files"]["LICENSE"]).decode().strip(), uri)
                restored = conversion.convert({"mode": "inspect", "files": result["files"]})
                root = next(item for item in restored["rocrate"]["@graph"] if item.get("@id") == "./")
                self.assertEqual(root["license"], {"@id": uri})

    def test_license_ambiguity(self):
        document = source()
        document["@graph"][1]["license"] = [{"@id": "https://example.org/license-a"}, {"@id": "https://example.org/license-b"}]
        with self.assertRaisesRegex(ValueError, "multiple licenses"):
            conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(document)})

    def test_context_aliases(self):
        document = source()
        document["@context"] = [document["@context"], {
            "name": "https://example.org/foreignName", "title": "http://schema.org/name",
            "about": "https://example.org/foreignAbout", "root": "http://schema.org/about",
        }]
        descriptor, dataset = document["@graph"][:2]
        descriptor["root"] = descriptor["about"]
        descriptor["about"] = {"@id": "#context"}
        dataset["title"] = dataset["name"]
        dataset["name"] = "Foreign property, not the title"
        original = json.dumps(document)
        result = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": original})
        restored = conversion.convert({"mode": "inspect", "files": result["files"]})
        root = next(item for item in restored["rocrate"]["@graph"] if item.get("@id") == "./")
        self.assertEqual(root["name"], "Supplied title")
        self.assertEqual(base64.b64decode(result["files"]["aruna-metadata.json"]).decode(), original)

    def test_scoped_context(self):
        document = source()
        document["@graph"][1]["@context"] = [
            {"name": "https://example.org/foreign", "title": "http://schema.org/name"},
            {"name": "http://schema.org/name", "title": "https://example.org/foreign"},
        ]
        document["@graph"][1]["title"] = "Foreign value"
        result = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(document)})
        restored = conversion.convert({"mode": "inspect", "files": result["files"]})
        root = next(item for item in restored["rocrate"]["@graph"] if item.get("@id") == "./")
        self.assertEqual(root["name"], "Supplied title")

    def test_literal_preserved(self):
        value = {"@id": "urn:root", "data": {"@value": {"@id": "urn:root"}, "@type": "@json"}}
        mapped = conversion.remap(value, "urn:root")
        self.assertEqual(mapped["@id"], "./")
        self.assertEqual(mapped["data"], value["data"])

    def test_rich_roundtrip(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scaffold(root)
            arc = ARC.load(directory)
            arc.Contacts = [Person(first_name="Ada", last_name="Example", email="ada@example.invalid")]
            table = arc.Assays[0].Tables[0]
            table.AddColumn(CompositeHeader.parameter(OntologyAnnotation("temperature", "PATO", "PATO:0000146")),
                            [CompositeCell.create_unitized("37", OntologyAnnotation("degree Celsius", "UO", "UO:0000027"))])
            original = arc.ToROCrateJsonString()
            generated = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": original})
            payload = "assays/assay/dataset/measurements.bin"
            self.assertIn(payload, generated["required"])
            self.assertNotIn(payload, generated["files"], "converter must not fabricate data bytes")
            generated["files"][payload] = base64.b64encode((root / payload).read_bytes()).decode()
            restored = conversion.convert({"mode": "inspect", "files": generated["files"]})["rocrate"]
            entities = restored["@graph"]
            self.assertTrue(any(item.get("additionalType") == "Study" and item.get("identifier") == "study" for item in entities))
            self.assertTrue(any(item.get("additionalType") == "Assay" and item.get("identifier") == "assay" for item in entities))
            self.assertTrue(any(item.get("@type") == "LabProcess" and item.get("object") and item.get("result") for item in entities))
            self.assertTrue(any(item.get("givenName") == "Ada" and item.get("familyName") == "Example" for item in entities))
            value = next(item for item in entities if item.get("@type") == "PropertyValue" and str(item.get("value")) == "37")
            self.assertEqual(value["unitText"], "degree Celsius")
            self.assertIn("0000027", value["unitCode"])
            self.assertIn("0000146", value["propertyID"])

    def test_sections_required(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scaffold(root)
            path = root / "isa.investigation.xlsx"
            book = load_workbook(path)
            sheet = book["isa_investigation"]
            row = next(index for index, values in enumerate(sheet.iter_rows(values_only=True), 1) if values[0] == "INVESTIGATION CONTACTS")
            sheet.delete_rows(row)
            book.save(path)
            book.close()
            with self.assertRaisesRegex(ValueError, "required sections"):
                conversion.convert({"mode": "inspect", "files": files(root)})

    def test_isa_authoritative(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scaffold(root)
            path = root / "isa.investigation.xlsx"
            book = load_workbook(path)
            sheet = book["isa_investigation"]
            row = next(index for index, values in enumerate(sheet.iter_rows(values_only=True), 1) if values[0] == "Investigation Title")
            sheet.cell(row, 2, "Changed ISA title")
            book.save(path)
            book.close()
            result = conversion.convert({"mode": "inspect", "files": files(root)})
            dataset = next(item for item in result["rocrate"]["@graph"] if item.get("@id") == "./")
            self.assertEqual(dataset["name"], "Changed ISA title")

    def test_registration_required(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scaffold(root)
            request = files(root)
            del request["studies/study/isa.study.xlsx"]
            with self.assertRaisesRegex(ValueError, "registered ISA study"):
                conversion.convert({"mode": "inspect", "files": request})

    def test_data_required(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scaffold(root)
            request = files(root)
            del request["assays/assay/dataset/measurements.bin"]
            with self.assertRaisesRegex(ValueError, "data is missing"):
                conversion.convert({"mode": "inspect", "files": request})

    def test_paths_confined(self):
        for path in ["../escape", "/tmp/escape", ".git/config", "folder/../../escape"]:
            with self.subTest(path=path), self.assertRaisesRegex(ValueError, "unsafe ARC path"):
                conversion.convert({"mode": "inspect", "files": {path: ""}})

    def test_external_entities(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory) / "arc"
            scaffold(root)
            private = Path(directory) / "outside.txt"
            private.write_text("outside marker")
            path = root / "isa.investigation.xlsx"
            with zipfile.ZipFile(path) as archive:
                entries = {name: archive.read(name) for name in archive.namelist()}
            name = "xl/worksheets/sheet1.xml"
            xml = entries[name]
            self.assertIn(b"Synthetic ARC experiment", xml)
            xml = xml.replace(b"Synthetic ARC experiment", b"&outside;")
            declaration = ('<!DOCTYPE worksheet [<!ENTITY outside SYSTEM "' + private.as_uri() + '">]>').encode()
            if xml.startswith(b"<?xml"):
                prefix, xml = xml.split(b"?>", 1)
                xml = prefix + b"?>" + declaration + xml
            else:
                xml = declaration + xml
            entries[name] = xml
            with zipfile.ZipFile(path, "w") as archive:
                for name, content in entries.items():
                    archive.writestr(name, content)
            with self.assertRaises(Exception):
                conversion.convert({"mode": "inspect", "files": files(root)})

    def test_facts_required(self):
        document = source()
        del document["@graph"][1]["description"]
        with self.assertRaisesRegex(ValueError, "description"):
            conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(document)})

    def test_cwl_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "workflow.cwl"
            path.write_text("cwlVersion: v1.2\nclass: CommandLineTool\nbaseCommand: echo\ninputs: []\noutputs: []\n")
            with patch("requests.sessions.Session.request", side_effect=AssertionError("network access forbidden")):
                conversion.cwl(root)
                path.write_text("cwlVersion: v1.2\nclass: CommandLineTool\ninputs: wrong\noutputs: []\n")
                with self.assertRaisesRegex(ValueError, "CWL v1.2"):
                    conversion.cwl(root)

    def test_cwl_confined(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "workflow.cwl"
            for target in ["https://example.invalid/secret.cwl", "file:///etc/passwd"]:
                path.write_text("$import: " + target + "\n")
                with patch("requests.sessions.Session.request", side_effect=AssertionError("network access forbidden")):
                    with self.subTest(target=target), self.assertRaises((ValueError, ValidationException)):
                        conversion.cwl(root)

    def edited(self, graph, change):
        generated = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(graph)})
        base = conversion.convert({"mode": "inspect", "files": generated["files"]})["rocrate"]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, content in generated["files"].items():
                (root / name).parent.mkdir(parents=True, exist_ok=True)
                (root / name).write_bytes(base64.b64decode(content))
            arc = ARC.load(directory)
            change(arc)
            arc.Write(str(root / "edited"))
            new = conversion.convert({"mode": "inspect", "files": files(root / "edited")})["rocrate"]
        return base, new

    def test_git_edits(self):
        graph = source()
        graph["@graph"][1].update({"creator": {"@id": "#person-ada"}, "keywords": "kept",
                                   "hasPart": [{"@id": "data/raw.csv"}]})
        graph["@graph"] += [{"@id": "#person-ada", "@type": "Person", "givenName": "Ada",
                             "familyName": "Lovelace", "email": "ada@example.org"},
                            {"@id": "data/raw.csv", "@type": "File", "name": "raw.csv"}]

        def change(arc):
            arc.Title = "Edited in ARCitect"
            arc.Contacts[0].EMail = "ada@example.com"
            study = ArcStudy.init("first")
            study.Title = "First study"
            arc.AddRegisteredStudy(study)

        base, new = self.edited(graph, change)
        merged = json.loads(conversion.convert({"mode": "merge", "graph": json.dumps(graph), "base": base, "new": new})["jsonld"])
        entities = {item["@id"]: item for item in merged["@graph"]}
        root = entities["urn:aruna:source"]
        self.assertEqual(merged["@context"][0], graph["@context"])
        self.assertEqual(root["name"], "Edited in ARCitect")
        self.assertEqual(root["keywords"], "kept")
        self.assertEqual(root["https://example.org/custom"], graph["@graph"][1]["https://example.org/custom"])
        self.assertEqual(root["creator"], {"@id": "#person-ada"})
        self.assertEqual(root["hasPart"], [{"@id": "data/raw.csv"}, {"@id": "studies/first/"}])
        self.assertEqual(entities["#person-ada"]["email"], "ada@example.com")
        self.assertEqual(entities["studies/first/"]["name"], "First study")
        self.assertNotIn("./", entities)
        self.assertEqual(entities["ro-crate-metadata.json"], graph["@graph"][0])
        repeated = conversion.convert({"mode": "merge", "graph": json.dumps(merged), "base": new, "new": new})
        self.assertIsNone(repeated["jsonld"])
        first = json.loads(conversion.convert({"mode": "merge", "graph": json.dumps(graph), "base": None, "new": new})["jsonld"])
        root = next(item for item in first["@graph"] if item["@id"] == "urn:aruna:source")
        self.assertEqual(root["name"], "Edited in ARCitect")
        self.assertNotIn("license", root)

    def test_assay_roundtrip(self):
        graph = source()
        generated = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(graph)})
        base = conversion.convert({"mode": "inspect", "files": generated["files"]})["rocrate"]
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            scaffold(root)
            pushed = {**generated["files"], **files(root)}
        new = conversion.convert({"mode": "inspect", "files": pushed})["rocrate"]
        merged = conversion.convert({"mode": "merge", "graph": json.dumps(graph), "base": base, "new": new})["jsonld"]
        terms = [entry for entry in json.loads(merged)["@context"] if isinstance(entry, dict)]
        self.assertEqual(terms[0]["LabProcess"], "https://bioschemas.org/LabProcess")
        regenerated = conversion.convert({"mode": "generate", "document_id": "document-id", "jsonld": merged})
        derived = json.loads(base64.b64decode(regenerated["files"]["ro-crate-metadata.json"]))
        self.assertEqual(sorted(map(json.dumps, derived["@graph"])), sorted(map(json.dumps, new["@graph"])))

    def test_git_removals(self):
        graph = source()
        graph["@graph"][1]["creator"] = [{"@id": "#person-ada"}, {"@id": "#person-bob"}]
        graph["@graph"] += [{"@id": "#person-ada", "@type": "Person", "givenName": "Ada", "familyName": "Lovelace"},
                            {"@id": "#person-bob", "@type": "Person", "givenName": "Bob", "familyName": "Builder"}]

        def change(arc):
            arc.Contacts.pop(1)

        base, new = self.edited(graph, change)
        merged = json.loads(conversion.convert({"mode": "merge", "graph": json.dumps(graph), "base": base, "new": new})["jsonld"])
        entities = {item["@id"]: item for item in merged["@graph"]}
        self.assertNotIn("#person-bob", entities)
        self.assertEqual(entities["urn:aruna:source"]["creator"], [{"@id": "#person-ada"}])
        self.assertEqual(entities["#context"], graph["@graph"][2])

    def test_metadata_edits(self):
        graph = source()
        edited = source()
        edited["@graph"][1]["https://example.org/custom"] = "edited in Git"
        edited["@graph"].append({"@id": "#new", "@type": "Thing", "name": "Added in Git"})
        current = source()
        current["@graph"][1]["name"] = "Concurrent graph edit"
        base = conversion.convert({"mode": "inspect", "files": conversion.convert(
            {"mode": "generate", "document_id": "document-id", "jsonld": json.dumps(graph)})["files"]})["rocrate"]
        merged = json.loads(conversion.convert({"mode": "merge", "graph": json.dumps(current), "base": base, "new": base,
                                                "json_base": json.dumps(graph), "json_new": json.dumps(edited)})["jsonld"])
        entities = {item["@id"]: item for item in merged["@graph"]}
        self.assertEqual(entities["urn:aruna:source"]["https://example.org/custom"], "edited in Git")
        self.assertEqual(entities["urn:aruna:source"]["name"], "Concurrent graph edit")
        self.assertEqual(entities["#new"]["name"], "Added in Git")
        unchanged = conversion.convert({"mode": "merge", "graph": json.dumps(current), "base": base, "new": base,
                                        "json_base": json.dumps(graph), "json_new": json.dumps(graph)})
        self.assertIsNone(unchanged["jsonld"])


if __name__ == "__main__":
    unittest.main()
