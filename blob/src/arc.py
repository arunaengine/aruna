"""Converts ISA RO-Crate metadata through the pinned DataPLANT ARCtrl library."""
# Copyright (c) 2026 The Aruna Contributors
# SPDX-License-Identifier: MIT or Apache-2.0

import base64
import importlib.metadata
import io
import json
import resource
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path, PurePosixPath
from urllib.parse import unquote, urlsplit

from arctrl import ARC
# ARCtrl 3.2.1 exposes graph JSON decoding through this generated extension name.
from arctrl.py.JsonIO.ldobject import ARCtrl_ROCrate_LDGraph__LDGraph_fromROCrateJsonString_Static_Z721C83C5 as read_graph
from arctrl.py.ROCrate.LDTypes.dataset import LDDataset
from openpyxl import load_workbook

LIMIT = 64 * 1024 * 1024


def confined(value):
    path = PurePosixPath(value)
    if (not value or path.is_absolute() or "\\" in value or "\0" in value
            or any(part in ("..", ".git") for part in path.parts)):
        raise ValueError("unsafe ARC path")
    return path


def remap(value, before):
    if isinstance(value, list):
        return [remap(item, before) for item in value]
    if isinstance(value, dict):
        if "@value" in value:
            return value
        return {key: "./" if key == "@id" and item == before else remap(item, before)
                for key, item in value.items()}
    return value


def adapt(value):
    if isinstance(value, list):
        return [adapt(item) for item in value]
    if not isinstance(value, dict) or "@value" in value:
        return value
    result = {}
    for key, item in value.items():
        if key == "@context":
            entries = item if isinstance(item, list) else [item]
            # ARCtrl searches arrays from the front; JSON-LD gives later contexts precedence.
            result[key] = ["https://w3id.org/ro/crate/1.2/context"
                           if entry == "https://w3id.org/ro/crate/1.3/context" else adapt(entry)
                           for entry in reversed(entries)]
        else:
            result[key] = adapt(item)
    return result


def prepare(source, identifier):
    document = adapt(json.loads(source))
    if not isinstance(document.get("@graph"), list):
        raise ValueError("RO-Crate @graph is required")
    graph = read_graph(json.dumps(document))
    context = graph.TryGetContext()
    descriptor = graph.TryGetNode("ro-crate-metadata.json")
    if descriptor is None:
        raise ValueError("RO-Crate metadata descriptor is required")
    root = descriptor.TryGetPropertyAsSingleNode("http://schema.org/about", graph, context)
    if root is None:
        raise ValueError("one RO-Crate root Dataset is required")
    if not LDDataset.try_get_identifier_as_string(root, context):
        entity = next(item for item in document["@graph"] if item.get("@id") == root.Id)
        entity["http://schema.org/identifier"] = identifier
    document = remap(document, root.Id)
    arc = ARC.from_rocrate_json_string(json.dumps(document))
    if not all(isinstance(value, str) and value.strip() for value in [arc.Title, arc.Description]):
        raise ValueError("ISA investigation requires a textual name and description")
    licenses = root.GetPropertyValues("http://schema.org/license", context=context)
    if len(licenses) > 1:
        raise ValueError("multiple licenses require an explicit ARC license document")
    license_value = licenses[0] if licenses else None
    if license_value is not None and (arc.License is None or not arc.License.Content):
        content = license_value if isinstance(license_value, str) else getattr(license_value, "Id", None)
        if not isinstance(content, str) or not content:
            raise ValueError("RO-Crate license cannot be represented")
        arc.SetLicenseFulltext(content)
    if arc.License is not None:
        arc.License.Path = "LICENSE"
    return arc


def workbook(path):
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > LIMIT:
            raise ValueError("expanded ISA workbook exceeds the limit")
    names = {"isa.investigation.xlsx": "isa_investigation", "isa.study.xlsx": "isa_study",
             "isa.assay.xlsx": "isa_assay"}
    sheet = names.get(path.name)
    if sheet:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet not in book.sheetnames:
                raise ValueError("ISA workbook is missing its metadata sheet")
            if path.name == "isa.investigation.xlsx":
                labels = {row[0] for row in book[sheet].iter_rows(values_only=True) if row}
                required = {"ONTOLOGY SOURCE REFERENCE", "INVESTIGATION", "INVESTIGATION PUBLICATIONS",
                            "INVESTIGATION CONTACTS", "Investigation Identifier", "Investigation Title",
                            "Investigation Description", "Investigation Submission Date",
                            "Investigation Public Release Date"}
                if not required <= labels:
                    raise ValueError("ISA investigation is missing required sections or labels")
        finally:
            book.close()


def data_paths(document):
    paths = []
    for entity in document["@graph"]:
        types = entity.get("@type", [])
        if "File" not in (types if isinstance(types, list) else [types]):
            continue
        value = entity["@id"]
        url = urlsplit(value)
        if url.scheme in ("http", "https"):
            continue
        if url.scheme or url.netloc:
            raise ValueError("unsupported ARC data reference")
        path = str(confined(unquote(url.path)))
        paths.append(path)
    return paths


def cwl(root):
    paths = list(root.rglob("*.cwl"))
    if not paths:
        return
    from cwltool.context import LoadingContext
    from cwltool.load_tool import fetch_document, resolve_and_validate_document
    from schema_salad.fetcher import DefaultFetcher

    class LocalFetcher(DefaultFetcher):
        def fetch_text(self, url, content_types=None):
            if isinstance(self.cache.get(url), str):
                return self.cache[url]
            parsed = urlsplit(url)
            if parsed.scheme != "file" or parsed.netloc:
                raise ValueError("CWL imports must stay inside the ARC")
            path = Path(unquote(parsed.path)).resolve()
            if not path.is_relative_to(root) or path.stat().st_size > LIMIT:
                raise ValueError("CWL import is outside the ARC or exceeds its limit")
            return super().fetch_text(url, content_types)

        def check_exists(self, url):
            parsed = urlsplit(url)
            if parsed.scheme != "file" or parsed.netloc:
                return url in self.cache
            path = Path(unquote(parsed.path)).resolve()
            return path.is_relative_to(root) and path.exists()

    for path in paths:
        try:
            context = LoadingContext({"fetcher_constructor": LocalFetcher, "do_update": False})
            context, document, uri = fetch_document(path.as_uri(), context)
            if document.get("cwlVersion") != "v1.2" or document.get("class") not in ("Workflow", "CommandLineTool"):
                raise ValueError("ARC workflows require CWL v1.2 Workflow or CommandLineTool")
            resolve_and_validate_document(context, document, uri)
        except Exception as error:
            raise ValueError("invalid or non-local CWL v1.2 description") from error


def inspect(root, require_data=True):
    if not (root / "isa.investigation.xlsx").is_file():
        raise ValueError("missing ISA investigation")
    for path in root.rglob("*.xlsx"):
        if path.name.startswith("isa."):
            workbook(path)
    cwl(root)
    arc = ARC.load(str(root))
    if not all(isinstance(value, str) and value.strip() for value in [arc.Identifier, arc.Title, arc.Description]):
        raise ValueError("ISA investigation identifier, title and description are required")
    studies = {item.Identifier: item for item in arc.Studies}
    assays = {item.Identifier: item for item in arc.Assays}
    for identifier in arc.RegisteredStudyIdentifiers:
        confined(identifier)
        if identifier not in studies or not (root / "studies" / identifier / "isa.study.xlsx").is_file():
            raise ValueError("registered ISA study is missing")
    for study in studies.values():
        for identifier in study.RegisteredAssayIdentifiers:
            confined(identifier)
            if identifier not in assays or not (root / "assays" / identifier / "isa.assay.xlsx").is_file():
                raise ValueError("registered ISA assay is missing")
    document = json.loads(arc.ToROCrateJsonString())
    root_entity = next(item for item in document["@graph"] if item.get("@id") == "./")
    # ARCtrl adds wall-clock export time, which cannot define a stable Git object.
    root_entity.pop("sdDatePublished", None)
    if not (root / "LICENSE").is_file() and root_entity.pop("license", None) is not None:
        # Without a LICENSE file ARCtrl reports its own default text, which nobody stated.
        document["@graph"] = [item for item in document["@graph"] if item.get("@id") != "#LICENSE"]
    elif arc.License is not None:
        content = arc.License.Content.strip()
        url = urlsplit(content)
        if url.scheme in ("http", "https") and url.netloc and not any(char.isspace() for char in content):
            root_entity["license"] = {"@id": content}
    if require_data and any(not (root / path).is_file() for path in data_paths(document)):
        raise ValueError("referenced ARC data is missing")
    return document


def canonical_workbook(path):
    output = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(output, "w") as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                properties = ET.fromstring(content)
                for field in ("created", "modified"):
                    value = properties.find("{http://purl.org/dc/terms/}" + field)
                    if value is not None:
                        value.text = "1980-01-01T00:00:00Z"
                content = ET.tostring(properties, encoding="utf-8")
            entry = zipfile.ZipInfo(name, (1980, 1, 1, 0, 0, 0))
            entry.create_system = 3
            entry.external_attr = 0o600 << 16
            target.writestr(entry, content)
    path.write_bytes(output.getvalue())


def generate(request, root):
    source = request["jsonld"]
    arc = prepare(source, request["document_id"])
    contracts = arc.GetWriteContracts()
    if len(contracts) > 10000:
        raise ValueError("generated ARC exceeds the path limit")
    for contract in contracts:
        confined(contract.path)
    arc.Write(str(root))
    document = inspect(root, require_data=False)
    for workbook in root.rglob("*.xlsx"):
        canonical_workbook(workbook)
    (root / "ro-crate-metadata.json").write_text(json.dumps(document, indent=2) + "\n")
    (root / "aruna-metadata.json").write_text(source)
    (root / ".gitattributes").write_text("*.bin filter=lfs diff=lfs merge=lfs -text\n")
    paths = sorted(path for path in root.rglob("*") if path.is_file())
    if len(paths) > 10000 or sum(path.stat().st_size for path in paths) > LIMIT // 2:
        raise ValueError("generated ARC exceeds the repository limit")
    return {"files": {str(path.relative_to(root)): base64.b64encode(path.read_bytes()).decode() for path in paths},
            "required": data_paths(document)}


def listed(value):
    return [] if value is None else value if isinstance(value, list) else [value]


def canon(value):
    return json.dumps(value, sort_keys=True)


def term(key):
    for prefix in ("http://schema.org/", "https://schema.org/", "schema:"):
        if key.startswith(prefix):
            return key[len(prefix):]
    return key


def entities(document):
    return {item["@id"]: item for item in (document or {}).get("@graph", [])
            if isinstance(item, dict) and isinstance(item.get("@id"), str)}


def properties(entity):
    return {term(key): listed(value) for key, value in (entity or {}).items() if key != "@id"}


def literals(entity):
    return {(key, canon(value)) for key, values in properties(entity).items() if key != "@type"
            for value in values if not (isinstance(value, dict) and "@id" in value)}


def identities(graph, root, wanted):
    """Maps ISA-derived ids to graph ids: same id, the root, or one best unique literal match."""
    mapping, claimed = {"./": root}, {root}
    for identifier, _ in wanted:
        if identifier in graph:
            mapping[identifier] = identifier
            claimed.add(identifier)
    candidates = {key: (set(map(str, listed(entity.get("@type")))), literals(entity))
                  for key, entity in graph.items()}
    for identifier, entity in wanted:
        if identifier in mapping:
            continue
        types, values = set(map(str, listed(entity.get("@type")))), literals(entity)
        scores = sorted(((len(values & known), key) for key, (kinds, known) in candidates.items()
                         if key not in claimed and types & kinds), reverse=True)
        if scores and scores[0][0] > 0 and (len(scores) == 1 or scores[1][0] < scores[0][0]):
            mapping[identifier] = scores[0][1]
            claimed.add(scores[0][1])
        else:
            mapping[identifier] = identifier
    return mapping


def apply(graph, root, base, new):
    """Applies the values changed from base to new onto graph; without base, new values win."""
    before, after = entities(base), entities(new)
    changed = [identifier for identifier in sorted(set(before) | set(after))
               if identifier != "ro-crate-metadata.json"
               and properties(before.get(identifier)) != properties(after.get(identifier))]
    referenced = {value["@id"] for identifier in changed
                  for entity in (before.get(identifier), after.get(identifier))
                  for values in properties(entity).values() for value in values
                  if isinstance(value, dict) and isinstance(value.get("@id"), str)}
    wanted = [(identifier, source[identifier]) for source in (before, after)
              for identifier in sorted(set(changed) | referenced) if identifier in source]
    mapping = identities(graph, root, wanted)

    def local(value):
        if isinstance(value, dict) and isinstance(value.get("@id"), str):
            return {**value, "@id": mapping.get(value["@id"], value["@id"])}
        return value

    removed = set()
    for identifier in changed:
        target, current = mapping[identifier], after.get(identifier)
        if current is None:
            if target != root and graph.pop(target, None) is not None:
                removed.add(target)
            continue
        entity = graph.setdefault(target, {"@id": target})
        old, current = properties(before.get(identifier)), properties(current)
        for name in sorted(set(old) | set(current)):
            key = next((key for key in entity if key != "@id" and term(key) == name), name)
            previous = {canon(local(value)) for value in old.get(name, [])}
            if base is None and name in current:
                previous = {canon(value) for value in listed(entity.get(key))}
            values = [local(value) for value in current.get(name, [])]
            if previous == {canon(value) for value in values}:
                continue
            kept = [value for value in listed(entity.get(key)) if canon(value) not in previous]
            for value in values:
                if canon(value) not in {canon(item) for item in kept}:
                    kept.append(value)
            if kept:
                entity[key] = kept if len(kept) > 1 or isinstance(entity.get(key), list) else kept[0]
            else:
                entity.pop(key, None)
    for entity in graph.values():
        for key in [key for key in entity if key != "@id"]:
            values = listed(entity[key])
            kept = [value for value in values if not (isinstance(value, dict) and value.get("@id") in removed)]
            if len(kept) != len(values):
                if kept:
                    entity[key] = kept if isinstance(entity[key], list) else kept[0]
                else:
                    del entity[key]


def merge(request):
    document = json.loads(request["graph"])
    graph = entities(document)
    about = graph.get("ro-crate-metadata.json", {}).get("about")
    root = about.get("@id") if isinstance(about, dict) else about
    if root not in graph:
        raise ValueError("RO-Crate root Dataset is required")
    before = canon(document)
    if None not in (request.get("json_base"), request.get("json_new")):
        apply(graph, root, json.loads(request["json_base"]), json.loads(request["json_new"]))
    apply(graph, root, request.get("base"), request["new"])
    document["@graph"] = list(graph.values())
    return {"jsonld": None if canon(document) == before else json.dumps(document)}


def convert(request):
    if request["mode"] == "merge":
        return merge(request)
    with tempfile.TemporaryDirectory(prefix="aruna-isa-") as directory:
        root = Path(directory)
        if request["mode"] == "generate":
            return generate(request, root)
        if request["mode"] != "inspect":
            raise ValueError("unsupported ARC conversion")
        files = request["files"]
        if len(files) > 10000:
            raise ValueError("ARC tree exceeds the path limit")
        for name, content in files.items():
            path = root / confined(name)
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(base64.b64decode(content, validate=True) if content else b"")
        return {"rocrate": inspect(root)}


if __name__ == "__main__":
    resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_CPU, (120, 120))
    if importlib.metadata.version("arctrl") != "3.2.1":
        raise RuntimeError("ARCtrl 3.2.1 is required")
    data = sys.stdin.buffer.read(LIMIT + 1)
    if len(data) > LIMIT:
        raise RuntimeError("ARC conversion input exceeds limit")
    try:
        result = convert(json.loads(data))
    except Exception:
        result = {"error": "ISA/ARC conversion failed: check required metadata, registrations and paths"}
    sys.stdout.write(json.dumps(result))
